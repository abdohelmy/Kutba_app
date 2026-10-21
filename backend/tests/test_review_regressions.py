import asyncio
from io import BytesIO
from threading import Event

import httpx
import pytest
from docx import Document

from app.api.routes import admin
from app.db.session import SessionLocal
from app.main import app
from app.models import Sermon, SermonSegment, SermonStatus
from app.services.canonical_sources import hadith_web_search_candidates
from app.services.documents import _extract_docx_text, split_text
from app.services.glossary import (
    build_glossary_entry,
    ensure_glossary_parentheticals,
    glossary_matches,
)
from app.services.retrieval import RetrievedSource
from app.services.sunnah_evidence import verify_sunnah_source


def test_multisentence_quotation_survives_segment_and_paragraph_boundaries():
    prefix = "عباد الله اتقوا الله وأحسنوا إلى الناس " * 65
    quotation = "«" + ". \n\n".join(["أحب لأخيك الخير وأحسن إلى جارك " * 10] * 3) + "»"
    text = prefix + quotation
    segments = split_text(text)
    assert any(quotation in segment for segment in segments)
    assert sum(len(hadith_web_search_candidates(s)) for s in segments) == 1
    assert " ".join(" ".join(segments).split()) == " ".join(text.split())


def test_quotation_longer_than_segment_limit_stays_searchable():
    quotation = "«" + "أحب لأخيك الخير وأحسن إلى جارك " * 10 + "»"
    segments = split_text("مقدمة قصيرة " + quotation + " ثم الخاتمة", max_chars=100)
    assert any(quotation in segment for segment in segments)
    assert sum(len(hadith_web_search_candidates(s)) for s in segments) == 1


def test_verse_and_its_reference_stay_together():
    verse = "قال تعالى: ﴿" + "الذين آمنوا وتطمئن قلوبهم بذكر الله " * 5 + "﴾ [الرعد: ٢٨]"
    text = "مقدمة قصيرة\n\n" + verse + "\n\nثم الخاتمة"
    assert any(verse in section for section in split_text(text, max_chars=100))


def test_long_multiline_ascii_hadith_is_never_split():
    quote = '"' + ("إنما الأعمال بالنيات ولكل امرئ ما نوى\n\n" * 150) + '"'
    text = "مقدمة\n\nقال رسول الله:\n\n" + quote + " رواه البخاري\n\nالخاتمة"
    sections = split_text(text, max_chars=100)
    assert any('قال رسول الله:\n\n' + quote + " رواه البخاري" in s for s in sections)


def test_unquoted_hadith_paragraph_is_never_split():
    paragraph = "قال رسول الله ﷺ " + "إنما الأعمال بالنيات ولكل امرئ ما نوى " * 100
    sections = split_text("مقدمة\n\n" + paragraph + "\n\nخاتمة", max_chars=100)
    assert paragraph.strip() in sections


def test_closing_formula_inside_nested_quote_does_not_split_it():
    quote = '«قال الرجل: "أقول قولي هذا وأستغفر الله" ثم جلس»'
    assert split_text(quote, max_chars=10) == [quote]


def test_reference_on_separate_paragraph_stays_with_verse():
    verse = "﴿الذين آمنوا وتطمئن قلوبهم بذكر الله﴾\n\n[الرعد: ٢٨]"
    assert split_text(verse, max_chars=10) == [verse]


def test_docx_preserves_table_and_nested_table_order():
    doc = Document()
    doc.add_paragraph("Opening")
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Quotation"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Nested quotation"
    cell.add_paragraph("Table ending")
    doc.add_paragraph("Closing")
    stream = BytesIO()
    doc.save(stream)
    assert _extract_docx_text(stream.getvalue()).split("\n\n") == [
        "Opening", "Quotation", "Nested quotation", "Table ending", "Closing"
    ]


def test_non_ascii_legacy_permission_password_does_not_bypass_approval(client):
    response = client.post("/api/v1/auth/register-mosque", json={
        "mosque_name": "Review mosque", "city": "Copenhagen", "country": "DK",
        "admin_display_name": "Admin", "username": "review", "password": "password",
        "permission_password": "٠١٠٨٢٠٢٥",
    })
    assert response.status_code == 201
    assert client.get("/api/v1/reader/mosques").json() == []


def test_blank_mosque_profile_is_rejected(client, seeded_accounts):
    from conftest import login_headers

    headers = login_headers(client, "admin", "correct-horse-123")
    assert client.patch("/api/v1/admin/profile", headers=headers,
                        json={"mosque_name": "   "}).status_code == 422
    assert client.get("/api/v1/admin/profile", headers=headers).json()["name"] == "Central Mosque"


def test_empty_sermon_cannot_be_published(client, seeded_accounts):
    from conftest import login_headers
    from test_workflow import _seed_sermon

    sermon_id = _seed_sermon(seeded_accounts)
    with SessionLocal() as db:
        sermon = db.get(Sermon, sermon_id)
        sermon.status = SermonStatus.REVIEW_REQUIRED
        sermon.segments.clear()
        db.commit()
    response = client.post(f"/api/v1/admin/sermons/{sermon_id}/publish",
                           headers=login_headers(client, "admin", "correct-horse-123"))
    assert response.status_code == 409


def test_older_untranslated_draft_is_resegmented_before_start(client, seeded_accounts, monkeypatch):
    from conftest import login_headers
    from test_workflow import _seed_sermon

    async def hold_job(sermon_id):
        pass

    monkeypatch.setattr(admin, "run_translation_job", hold_job)
    sermon_id = _seed_sermon(seeded_accounts)
    with SessionLocal() as db:
        sermon = db.get(Sermon, sermon_id)
        sermon.arabic_text = "قال الخطيب: «إنما الأعمال بالنيات ولكل امرئ ما نوى»"
        sermon.segments[0].arabic_text = "قال الخطيب: «إنما الأعمال بالنيات"
        sermon.segments.append(SermonSegment(ordinal=1, arabic_text="ولكل امرئ ما نوى»"))
        db.commit()
    response = client.post(f"/api/v1/admin/sermons/{sermon_id}/translate",
                           headers=login_headers(client, "admin", "correct-horse-123"))
    assert response.status_code == 200
    with SessionLocal() as db:
        sermon = db.get(Sermon, sermon_id)
        assert len(sermon.segments) == 1
        assert sermon.segments[0].arabic_text == sermon.arabic_text


@pytest.mark.asyncio
async def test_upload_does_not_block_health_requests(client, seeded_accounts, monkeypatch):
    from conftest import login_headers

    headers = login_headers(client, "admin", "correct-horse-123")
    entered, release = Event(), Event()

    def extract(*args, **kwargs):
        entered.set()
        if not release.wait(3):
            raise RuntimeError("The event loop could not release the extraction worker")
        return "الحمد لله رب العالمين والصلاة والسلام على رسول الله وعلى آله وصحبه أجمعين"

    monkeypatch.setattr(admin, "extract_document_text", extract)
    doc = Document()
    doc.add_paragraph("Test document")
    stream = BytesIO()
    doc.save(stream)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as api:
        upload = asyncio.create_task(api.post("/api/v1/admin/sermons", headers=headers,
            data={"khutba_date": "2026-09-25", "target_language": "en"},
            files={"file": ("test.docx", stream.getvalue(), "application/octet-stream")}))
        try:
            assert await asyncio.to_thread(entered.wait, 2)
            assert (await asyncio.wait_for(api.get("/health"), 1)).status_code == 200
        finally:
            release.set()
        assert (await upload).status_code == 201


def test_unmarked_taqwa_gets_complete_highlight_and_arabic():
    arabic = "التقوى تحفظ المؤمن"
    english = "God-consciousness protects the believer."
    matches = glossary_matches(arabic, english)
    match = next(item for item in matches if item.entry.arabic_term == "التقوى")
    assert match.display_term == "God-consciousness"
    assert english[match.translation_start:match.translation_end] == "God-consciousness"
    repaired, issues = ensure_glossary_parentheticals(arabic, english, [])
    # A separate ordinary sense now defines believer without reusing Allah's name.
    assert repaired == "God-consciousness (التقوى) protects the believer (المؤمن)."
    assert not issues


def test_fallback_does_not_guess_between_identical_renderings():
    entries = (
        build_glossary_entry("العين", "An eye used for sight.", "eye"),
        build_glossary_entry("العين", "A spring or source of water.", "eye"),
    )
    assert not glossary_matches("العين", "The eye.", entries=entries)


def test_glossary_repair_preserves_canonical_wording():
    english = "God-consciousness protects the believer."
    repaired, _ = ensure_glossary_parentheticals("التقوى تحفظ المؤمن", english, [],
                                               protected_spans=[(0, len(english))])
    assert repaired == english
    assert glossary_matches("التقوى تحفظ المؤمن", repaired)


def _source():
    return RetrievedSource(chunk_id="canonical:sunnah:bukhari:1", source_id="sunnah.com:bukhari:1",
        title="Sahih al-Bukhari 1", authority="Sunnah.com", score=10000,
        text="Actions are judged by intentions.", arabic_text="إنما الأعمال بالنيات",
        source_kind="hadith", url="https://sunnah.com/bukhari:1")


@pytest.mark.asyncio
@pytest.mark.parametrize("page,valid", [
    ("<p>إنما الأعمال بالنيات</p><p>Actions are <b>judged</b> by intentions.</p>", True),
    ("<p>إنما الأعمال بالنيات</p><p>Unrelated English.</p>", False),
    ("<p>الدين النصيحة</p><p>Actions are judged by intentions.</p>", False),
    ("<script>إنما الأعمال بالنيات Actions are judged by intentions.</script>", False),
])
async def test_hadith_page_must_support_both_languages(page, valid):
    async with httpx.AsyncClient(transport=httpx.MockTransport(
        lambda request: httpx.Response(200, text=page)
    )) as client:
        issue = await verify_sunnah_source(_source(), client)
    assert (issue is None) == valid


@pytest.mark.asyncio
async def test_unavailable_hadith_evidence_is_a_review_issue():
    async with httpx.AsyncClient(transport=httpx.MockTransport(
        lambda request: httpx.Response(403)
    )) as client:
        issue = await verify_sunnah_source(_source(), client)
    assert "unavailable" in issue


@pytest.mark.asyncio
async def test_hadith_evidence_does_not_follow_external_redirect():
    requests = []

    def respond(request):
        requests.append(str(request.url))
        return httpx.Response(302, headers={"location": "https://example.org/hadith"})

    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        issue = await verify_sunnah_source(_source(), client)
    assert "left Sunnah.com" in issue
    assert requests == ["https://sunnah.com/bukhari:1"]
