from io import BytesIO

from docx import Document
from pypdf import PdfWriter

from app.services.documents import DocumentKind, UploadedDocument, extract_document_text, split_text


def test_split_text_keeps_order_and_bounds():
    text = "\n\n".join(["الحمد لله " * 20, "أما بعد " * 20, "فاتقوا الله " * 20])
    segments = split_text(text, max_chars=150)
    assert len(segments) >= 3
    assert all(len(segment) <= 150 for segment in segments)
    assert segments[0].startswith("الحمد")


def test_first_khutba_closing_formula_forces_a_segment_boundary():
    text = (
        "الحمد لله ونذكر موضوع الخطبة. أقول قولي هذا وأستغفر الله العظيم لي ولكم. "
        "ثم قام الخطيب للخطبة الثانية وحمد الله."
    )

    segments = split_text(text, max_chars=500)

    assert len(segments) == 2
    assert "أستغفر الله العظيم لي ولكم" in segments[0]
    assert segments[1].startswith("ثم قام الخطيب")


def test_long_segmentation_does_not_cut_inside_a_bounded_quotation():
    quotation = "إنما الأعمال بالنيات ولكل امرئ ما نوى وهذا نص متصل"
    text = f"{'مقدمة ' * 12} «{quotation}» {'خاتمة ' * 12}"

    segments = split_text(text, max_chars=90)

    assert all(len(segment) <= 90 for segment in segments)
    assert any(quotation in segment for segment in segments)


def test_extracts_arabic_text_from_docx():
    output = BytesIO()
    document = Document()
    document.add_heading("خطبة الجمعة", level=1)
    document.add_paragraph("الحمد لله رب العالمين، والصلاة والسلام على رسول الله.")
    document.add_paragraph("أما بعد، فاتقوا الله عباد الله، واصبروا واشكروا.")
    document.save(output)

    text = extract_document_text(
        UploadedDocument(output.getvalue(), DocumentKind.DOCX),
        require_arabic=True,
    )

    assert "الحمد لله رب العالمين" in text
    assert "فاتقوا الله" in text


def test_scanned_pdf_uses_ocr_fallback(monkeypatch):
    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(output)
    ocr_text = (
        "الحمد لله رب العالمين، نحمده ونستعينه ونستغفره. "
        "أما بعد، فاتقوا الله عباد الله واصبروا واشكروا."
    )
    monkeypatch.setattr("app.services.documents.ocr_pdf", lambda content, page_count: ocr_text)

    text = extract_document_text(
        UploadedDocument(output.getvalue(), DocumentKind.PDF),
        require_arabic=True,
        allow_ocr=True,
    )

    assert text == ocr_text
