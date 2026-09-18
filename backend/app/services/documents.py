import hashlib
import re
from dataclasses import dataclass
from enum import StrEnum
from io import BytesIO
from pathlib import Path, PurePosixPath
from uuid import uuid4
from zipfile import BadZipFile, ZipFile

from docx import Document
from fastapi import HTTPException, UploadFile, status
from pypdf import PdfReader

from app.core.config import get_settings
from app.services.ocr import ocr_pdf

ARABIC_RE = re.compile(r"[\u0600-\u06FF]")
ARABIC_MARKS_RE = re.compile(r"[\u0610-\u061A\u064B-\u065F\u0670\u06D6-\u06ED]")
KHUTBA_CLOSE_RE = re.compile(
    r"اقول\s+قولي\s+هذا\s+واستغفر\s+الله(?:\s+العظيم)?"
    r"(?:\s+لي\s+ولكم(?:\s+ولسائر\s+المسلمين)?)?"
)
PDF_MIME_TYPES = {"application/pdf", "application/x-pdf"}
DOCX_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/zip",
}
GENERIC_MIME_TYPES = {"", "application/octet-stream"}
MAX_DOCX_ENTRIES = 2_000
MAX_DOCX_UNCOMPRESSED_BYTES = 30 * 1024 * 1024


class DocumentKind(StrEnum):
    PDF = "pdf"
    DOCX = "docx"


@dataclass(frozen=True)
class UploadedDocument:
    content: bytes
    kind: DocumentKind


def _inspect_docx(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise ValueError("The upload is not a valid DOCX document")
            if len(entries) > MAX_DOCX_ENTRIES:
                raise ValueError("The DOCX document contains too many embedded files")
            if sum(entry.file_size for entry in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ValueError("The expanded DOCX document is too large")
            for entry in entries:
                path = PurePosixPath(entry.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise ValueError("The DOCX document contains an unsafe file path")
    except BadZipFile as exc:
        raise ValueError("The upload is not a valid DOCX document") from exc


def _detect_document_kind(content: bytes) -> DocumentKind:
    if content.startswith(b"%PDF"):
        return DocumentKind.PDF
    if content.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
        _inspect_docx(content)
        return DocumentKind.DOCX
    raise ValueError("Only valid PDF or DOCX documents are accepted")


async def read_document_upload(upload: UploadFile) -> UploadedDocument:
    limit = get_settings().max_document_bytes
    content = await upload.read(limit + 1)
    if len(content) > limit:
        raise HTTPException(
            status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            "The document is too large",
        )
    if not content:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "The uploaded document is empty")

    try:
        kind = _detect_document_kind(content)
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc)) from exc

    content_type = (upload.content_type or "").lower()
    allowed_mime_types = PDF_MIME_TYPES if kind == DocumentKind.PDF else DOCX_MIME_TYPES
    if content_type not in allowed_mime_types | GENERIC_MIME_TYPES:
        raise HTTPException(
            status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            f"The declared file type does not match the uploaded {kind.value.upper()} document",
        )

    suffix = Path(upload.filename or "").suffix.lower()
    if suffix and suffix != f".{kind.value}":
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"The filename extension does not match the uploaded {kind.value.upper()} document",
        )
    return UploadedDocument(content=content, kind=kind)


def _extract_pdf_text(content: bytes) -> tuple[str, int]:
    try:
        reader = PdfReader(BytesIO(content))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
    except Exception as exc:
        raise ValueError("The PDF could not be read") from exc
    return "\n\n".join(page for page in pages if page).strip(), len(reader.pages)


def _extract_docx_text(content: bytes) -> str:
    try:
        document = Document(BytesIO(content))
        blocks = [
            paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()
        ]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append(" | ".join(cells))
    except Exception as exc:
        raise ValueError("The DOCX document could not be read") from exc
    return "\n\n".join(blocks).strip()


def extract_document_text(
    document: UploadedDocument,
    require_arabic: bool = False,
    allow_ocr: bool = False,
) -> str:
    if document.kind == DocumentKind.PDF:
        text, page_count = _extract_pdf_text(document.content)
        if len(text) < 40 and allow_ocr:
            text = ocr_pdf(document.content, page_count)
        empty_message = (
            "No usable text was found. Scanned/image-only PDFs need an OCR service before upload."
        )
    else:
        text = _extract_docx_text(document.content)
        empty_message = "No usable text was found in the DOCX document."

    if len(text) < 40:
        raise ValueError(empty_message)
    if require_arabic:
        validate_arabic_text(text)
    return text


def validate_arabic_text(text: str) -> None:
    arabic_count = len(ARABIC_RE.findall(text))
    if arabic_count < 20 or arabic_count / max(len(text), 1) < 0.10:
        raise ValueError("The extracted sermon does not appear to contain enough Arabic text")


def store_document(document: UploadedDocument, category: str) -> tuple[str, str]:
    settings = get_settings()
    folder = settings.storage_dir / category
    folder.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(document.content).hexdigest()
    path = folder / f"{uuid4()}.{document.kind.value}"
    path.write_bytes(document.content)
    return str(path), digest


def remove_stored_file(path: str) -> None:
    file_path = Path(path)
    storage_root = get_settings().storage_dir.resolve()
    try:
        resolved = file_path.resolve()
        if resolved.is_relative_to(storage_root) and resolved.is_file():
            resolved.unlink()
    except (OSError, ValueError):
        return


def _normalized_arabic_with_offsets(value: str) -> tuple[str, list[int]]:
    characters: list[str] = []
    offsets: list[int] = []
    character_translation = {
        "أ": "ا",
        "إ": "ا",
        "آ": "ا",
        "ٱ": "ا",
        "ى": "ي",
        "ؤ": "و",
        "ئ": "ي",
    }
    for index, character in enumerate(value):
        if ARABIC_MARKS_RE.fullmatch(character) or character == "ـ":
            continue
        characters.append(character_translation.get(character, character))
        offsets.append(index)
    return "".join(characters), offsets


def _khutba_parts(text: str) -> list[str]:
    normalized, offsets = _normalized_arabic_with_offsets(text)
    boundaries = [
        offsets[match.end() - 1] + 1
        for match in KHUTBA_CLOSE_RE.finditer(normalized)
        if match.end() > 0
    ]
    if not boundaries:
        return [text]
    parts: list[str] = []
    start = 0
    for end in boundaries:
        while end < len(text) and text[end] in " \t،,.;؛!?؟":
            end += 1
        part = text[start:end].strip()
        if part:
            parts.append(part)
        start = end
    remainder = text[start:].strip()
    if remainder:
        parts.append(remainder)
    return parts


def _safe_long_pieces(value: str, max_chars: int) -> list[str]:
    from app.services.canonical_sources import bounded_passage_candidates

    protected = [
        (passage.start, passage.end)
        for passage in bounded_passage_candidates(value)
        if passage.start is not None and passage.end is not None
    ]
    pieces: list[str] = []
    start = 0
    while len(value) - start > max_chars:
        target = start + max_chars
        candidates = [
            match.end()
            for match in re.finditer(r"\s+", value[start:target])
            if not any(
                protected_start < start + match.end() < protected_end
                for protected_start, protected_end in protected
            )
        ]
        end = start + candidates[-1] if candidates else target
        piece = value[start:end].strip()
        if piece:
            pieces.append(piece)
        start = end
        while start < len(value) and value[start].isspace():
            start += 1
    remainder = value[start:].strip()
    if remainder:
        pieces.append(remainder)
    return pieces


def split_text(text: str, max_chars: int = 2800) -> list[str]:
    segments: list[str] = []
    current = ""
    khutba_parts = _khutba_parts(text)
    for part_index, khutba_part in enumerate(khutba_parts):
        paragraphs = [
            re.sub(r"[ \t]+", " ", value).strip() for value in re.split(r"\n\s*\n", khutba_part)
        ]
        for paragraph in (value for value in paragraphs if value):
            sentences = (
                [item.strip() for item in re.split(r"(?<=[.!؟؛])\s+", paragraph) if item.strip()]
                if len(paragraph) > max_chars
                else [paragraph]
            )
            for sentence in sentences:
                for piece in _safe_long_pieces(sentence, max_chars):
                    candidate = f"{current}\n\n{piece}".strip() if current else piece
                    if current and len(candidate) > max_chars:
                        segments.append(current)
                        current = piece
                    else:
                        current = candidate
        if part_index < len(khutba_parts) - 1 and current:
            segments.append(current)
            current = ""
    if current:
        segments.append(current)
    return segments
